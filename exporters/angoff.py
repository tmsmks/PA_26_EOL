"""Grille Excel des seuils de niveau (Angoff modifié), à remplir par les juges."""
from __future__ import annotations

import io

from exporters.tables import ILLEGAL_XML_CHARS
from services.angoff_service import MIN_JUDGES, AngoffPoint

COLUMNS = [
    ("point", 14), ("chapitre", 8), ("scène", 26), ("personnage", 16), ("réplique", 55),
    ("réponses proposées", 80), ("réponse la plus probable (n°)", 14), ("commentaire", 30),
]


def _clean(text):
    return ILLEGAL_XML_CHARS.sub("", text) if isinstance(text, str) else text


def angoff_workbook(book_title: str, points: list[AngoffPoint], scope: str, year: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    to_fill = PatternFill("solid", fgColor="FFF7D6")

    wb = Workbook()
    info = wb.active
    info.title = "consignes"
    info.column_dimensions["A"].width = 32
    info.column_dimensions["B"].width = 90
    lines = [
        (f"Seuils de niveau — {book_title}", None),
        (f"{scope} · {len(points)} points de choix", None),
        (None, None),
        ("Votre nom ou code de juge", ""),
        (None, None),
        ("L'étudiant·e à imaginer", f"Public du livre : {year}. Imaginez un·e étudiant·e de ce public, "
         "« tout juste compétent·e » : ses réponses sont acceptables pour son année, sans plus. Il ou elle "
         "n'est ni brillant·e ni en difficulté."),
        ("Ce qu'on vous demande", "Pour chaque point de choix de la feuille « seuils », lisez la réplique et "
         "les réponses proposées, puis écrivez le numéro de la réponse que cet·te étudiant·e choisirait le "
         "plus probablement. Ce n'est pas forcément la meilleure réponse."),
        ("À quoi cela sert", "Vos choix donnent, chapitre par chapitre, la note de cet·te étudiant·e "
         "imaginaire : c'est le seuil de la bande « Compétent ». La médiane des juges est appliquée à partir "
         f"de {MIN_JUDGES} juges (méthode d'Angoff modifiée ; Dwyer et al. 2016)."),
        ("À l'aveugle", "Les points des réponses ne figurent pas dans la grille. Ne vous concertez pas "
         "avant d'avoir rendu votre grille."),
        ("Retour", "Enregistrez le fichier rempli (.xlsx) et renvoyez-le à la personne qui coordonne."),
    ]
    for label, value in lines:
        info.append([_clean(label), _clean(value)])
    for row in info.iter_rows(min_row=1, max_row=info.max_row):
        row[0].font = bold
        for cell in row:
            cell.alignment = wrap
    info["A1"].font = Font(bold=True, size=14)
    info["B4"].fill = to_fill
    wb.defined_names["juge"] = DefinedName("juge", attr_text="consignes!$B$4")

    ws = wb.create_sheet("seuils")
    ws.append([name for name, _ in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        ws.cell(row=1, column=i).font = bold
    for p in points:
        listed = "\n".join(f"{i}. {t}" for i, t in enumerate(p.responses, start=1))
        ws.append([_clean(v) for v in (p.ref, p.chapter_id, p.scene, p.actor, p.prompt, listed, None, None)])
    n = len(points) + 1
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row:
            cell.alignment = wrap
        row[6].fill = to_fill
        row[7].fill = to_fill
    ws.freeze_panes = "A2"
    pick = DataValidation(type="whole", operator="between", formula1="1", formula2="6", allow_blank=True,
                          showErrorMessage=True, errorTitle="Valeur refusée", error="Un numéro de réponse.")
    pick.add(f"G2:G{n}")
    ws.add_data_validation(pick)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
