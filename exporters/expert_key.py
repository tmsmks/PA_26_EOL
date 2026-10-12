"""Classeurs Excel de la validation de la clé de cotation.

- `grid_workbook` : la grille qu'un expert remplit (consignes, grille ancrée,
  une ligne par réponse, cases à saisie contrôlée de −3 à +3).
- `synthesis_tables` : la synthèse des cotes reçues, écrite avec les tables
  d'export (`exporters.tables.to_xlsx`).
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from domain.expert_rubric import CRITICAL_DEFINITION, RUBRIC, RUBRIC_STATUS
from domain.soft_skills import SOFT_SKILL_LABELS, SOFT_SKILLS
from exporters.tables import ILLEGAL_XML_CHARS, Table
from services.expert_key_service import SCORE_MAX, SCORE_MIN, KeyItem

GRID_COLUMNS = [
    ("ref", 14), ("chapitre", 8), ("scène", 26), ("personnage", 16), ("réplique", 60),
    ("réponse n°", 9), ("réponse", 60),
    *((SOFT_SKILL_LABELS[k], 13) for k in SOFT_SKILLS),
    ("critique", 10), ("justification", 50), ("commentaire", 30),
]


def _clean(text: Any) -> Any:
    return ILLEGAL_XML_CHARS.sub("", text) if isinstance(text, str) else text


def grid_workbook(book_title: str, items: list[KeyItem], round_: int, scope: str) -> bytes:
    """Grille à l'aveugle : les points de l'auteur n'y figurent pas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    to_fill = PatternFill("solid", fgColor="FFF7D6")

    wb = Workbook()
    info = wb.active
    info.title = "consignes"
    info.column_dimensions["A"].width = 30
    for col in "BCDE":
        info.column_dimensions[col].width = 45
    lines = [
        (f"Validation de la clé de cotation — {book_title}", None),
        (f"{scope} · {len(items)} réponses · tour {round_}", None),
        (None, None),
        ("Votre nom ou code d'expert", ""),
        ("Tour", round_),
        (None, None),
        ("Ce qu'on vous demande", "Pour chaque réponse de la feuille « cotation », notez de −3 à +3 "
         "ce qu'elle montre de chacune des 7 compétences, selon la grille ci-dessous. Jugez la "
         "réponse dans sa situation (la réplique qui précède), comme le ferait un·e étudiant·e "
         "de 2ᵉ année."),
        ("Case vide", "Une compétence laissée vide sur une réponse cotée vaut 0 : neutre, ou pas "
         "en jeu. Une ligne entièrement vide est considérée comme non cotée."),
        ("Critique", CRITICAL_DEFINITION + " Écrivez « oui » dans la colonne critique."),
        ("Justification", "Une phrase par réponse, si possible : elle sera montrée aux étudiant·e·s "
         "dans le feedback. Utilisez « commentaire » pour une remarque sur la réponse elle-même "
         "(formulation, réalisme)."),
        ("Cotation à l'aveugle", "Les points donnés par l'auteur ne figurent pas dans cette grille. "
         "Ne vous concertez pas avant d'avoir rendu votre grille."),
        ("Retour", "Enregistrez le fichier rempli (.xlsx) et renvoyez-le à la personne qui "
         "coordonne la validation."),
        (None, None),
        ("Grille ancrée", RUBRIC_STATUS),
    ]
    for label, value in lines:
        info.append([_clean(label), _clean(value)])
    for row in info.iter_rows(min_row=1, max_row=info.max_row):
        row[0].font = bold
        for cell in row:
            cell.alignment = wrap
    info["A1"].font = Font(bold=True, size=14)
    info["B4"].fill = to_fill
    info["B5"].fill = to_fill
    wb.defined_names["expert"] = DefinedName("expert", attr_text="consignes!$B$4")
    wb.defined_names["tour"] = DefinedName("tour", attr_text="consignes!$B$5")

    info.append(["Compétence", "Ancrage", f"{SCORE_MIN}", "0", f"+{SCORE_MAX}"])
    for cell in info[info.max_row]:
        cell.font = bold
    for k in SOFT_SKILLS:
        r = RUBRIC[k]
        info.append([SOFT_SKILL_LABELS[k], r["anchor"], r["-3"], r["0"], r["+3"]])
        for cell in info[info.max_row]:
            cell.alignment = wrap

    ws = wb.create_sheet("cotation")
    ws.append([name for name, _ in GRID_COLUMNS])
    for i, (_, width) in enumerate(GRID_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        ws.cell(row=1, column=i).font = bold
    for it in items:
        ws.append([_clean(v) for v in (
            it.ref, it.chapter_id, it.scene, it.actor, it.prompt, it.response_no, it.text,
            *([None] * len(SOFT_SKILLS)), None, None, None,
        )])
    first_score = 8  # colonne H
    last_score = first_score + len(SOFT_SKILLS) - 1
    n = len(items) + 1
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row:
            cell.alignment = wrap
        for cell in row[first_score - 1:]:
            cell.fill = to_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(GRID_COLUMNS))}{n}"

    scores = DataValidation(
        type="whole", operator="between", formula1=str(SCORE_MIN), formula2=str(SCORE_MAX),
        allow_blank=True, showErrorMessage=True,
        errorTitle="Valeur refusée", error=f"Un entier de {SCORE_MIN} à +{SCORE_MAX}, ou vide.",
    )
    scores.add(f"{get_column_letter(first_score)}2:{get_column_letter(last_score)}{n}")
    critical = DataValidation(type="list", formula1='"oui,non"', allow_blank=True)
    critical.add(f"{get_column_letter(last_score + 1)}2:{get_column_letter(last_score + 1)}{n}")
    ws.add_data_validation(scores)
    ws.add_data_validation(critical)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def synthesis_tables(analysis: dict[str, Any], round_: int) -> list[Table]:
    """Accord par compétence, puis une ligne par réponse cotée (pour le tour suivant)."""
    agreement = Table(
        "accord",
        "Accord entre experts par compétence (réponses cotées par tous les experts du tour).",
        [
            ("competence", "Compétence, ou total des 7 compétences."),
            ("reponses", "Réponses cotées par tous les experts du tour."),
            ("experts", "Nombre d'experts."),
            ("icc_un_expert", "ICC(2,1) : fiabilité de la cote d'un seul expert."),
            ("niveau_un_expert", "Interprétation de Koo & Li (2016)."),
            ("icc_moyenne", "ICC(2,k) : fiabilité de la moyenne des k experts."),
            ("niveau_moyenne", "Interprétation de Koo & Li (2016)."),
        ],
        [
            {
                "competence": a["label"],
                "reponses": a["n_items"],
                "experts": len(analysis["experts"]),
                "icc_un_expert": a["icc_single"],
                "niveau_un_expert": a["level_single"],
                "icc_moyenne": a["icc_average"],
                "niveau_moyenne": a["level_average"],
            }
            for a in analysis["agreement"]
        ],
    )

    columns = [
        ("ref", "Identifiant de la réponse (chapitre, scène, interaction, rang)."),
        ("chapitre", "Numéro du chapitre."),
        ("scene", "Titre de la scène."),
        ("personnage", "Personnage qui parle."),
        ("replique", "Réplique qui précède la réponse."),
        ("reponse", "Texte de la réponse."),
        ("experts", "Experts qui ont coté cette réponse."),
        ("a_revoir", "1 si les experts divergent (écart ≥ 3) ou s'écartent de l'auteur (≥ 2)."),
        ("divergences", "Compétences où les experts divergent."),
        ("ecarts_auteur", "Compétences où la médiane s'écarte de la clé de l'auteur."),
    ]
    for k in SOFT_SKILLS:
        label = SOFT_SKILL_LABELS[k]
        columns += [
            (f"{label} auteur", f"{label} : points de l'auteur (clé provisoire)."),
            (f"{label} médiane", f"{label} : médiane des experts."),
            (f"{label} min", f"{label} : cote la plus basse."),
            (f"{label} max", f"{label} : cote la plus haute."),
        ]
    columns += [
        ("total auteur", "Total des 7 compétences, clé de l'auteur."),
        ("total médian", "Médiane des totaux des experts."),
        ("critique votes", "Experts qui la jugent critique."),
        ("critique majorite", "1 si plus de la moitié des experts la jugent critique."),
        ("critique provisoire", "1 si l'auteur l'a marquée critique, ou si son total est ≤ −7."),
        ("justifications", "Justifications des experts."),
        ("commentaires", "Commentaires des experts."),
    ]
    rows = []
    for r in analysis["rows"]:
        it: KeyItem = r["item"]
        row = {
            "ref": it.ref,
            "chapitre": it.chapter_id,
            "scene": it.scene,
            "personnage": it.actor,
            "replique": it.prompt,
            "reponse": it.text,
            "experts": r["n_experts"],
            "a_revoir": r["to_review"],
            "divergences": ", ".join(SOFT_SKILL_LABELS[k] for k in r["disagreement"]),
            "ecarts_auteur": ", ".join(SOFT_SKILL_LABELS[k] for k in r["author_gap"]),
            "total auteur": it.author_total,
            "total médian": r["median_total"],
            "critique votes": r["critical_votes"],
            "critique majorite": r["critical_majority"],
            "critique provisoire": it.provisional_critical,
            "justifications": " | ".join(r["rationales"]),
            "commentaires": " | ".join(r["comments"]),
        }
        for k in SOFT_SKILLS:
            label = SOFT_SKILL_LABELS[k]
            st = r["stats"][k]
            row |= {
                f"{label} auteur": it.author[k],
                f"{label} médiane": st["median"],
                f"{label} min": st["min"],
                f"{label} max": st["max"],
            }
        rows.append(row)
    responses = Table("reponses", f"Une réponse cotée par ligne (tour {round_}).", columns, rows)
    return [agreement, responses]


def synthesis_meta(book_title: str, analysis: dict[str, Any], round_: int) -> dict[str, Any]:
    return {
        "livre": book_title,
        "tour": round_,
        "exporte_le": datetime.now(),
        "experts": ", ".join(analysis["experts"]) or "aucun",
        "reponses_cotees": analysis["n_rated"],
        "reponses_a_revoir": analysis["n_to_review"],
        "remarques": [
            "Cotes de −3 à +3 par compétence ; une compétence laissée vide sur une réponse cotée vaut 0.",
            "ICC à deux facteurs, effets aléatoires, accord absolu. Formes et seuils de Koo & Li "
            "(2016) : < 0,5 faible, 0,5–0,75 modéré, 0,75–0,9 bon, ≥ 0,9 excellent.",
            "À revoir au tour suivant (Delphi ; Jünger et al. 2017) : les réponses où les "
            f"experts s'écartent de 3 points ou plus, ou dont la médiane s'écarte de l'auteur de 2 "
            "points ou plus.",
            "La clé consensuelle (médiane) n'est pas encore appliquée au livre : les points du jeu "
            "restent ceux de l'auteur.",
        ],
    }
