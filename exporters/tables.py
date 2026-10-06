"""Tables de données → JSON, CSV (archive zip) ou Excel (.xlsx).

Un export est une liste de `Table` (nom, colonnes décrites, lignes) et un
dictionnaire `meta` (date, périmètre, remarques). Les trois formats portent
les mêmes tables et le même dictionnaire des colonnes :
- JSON : {"meta", "dictionnaire", "tables"}, dates ISO 8601 avec fuseau ;
- CSV : un fichier par table dans une archive zip, UTF-8 avec BOM (accents
  lisibles dans Excel), séparateur virgule, point décimal, booléens 1/0 ;
- Excel : une feuille par table, plus une feuille « lisez-moi ».

Un texte qui commence comme une formule (« = », « + », « - », « @ ») n'est
jamais exécuté par le tableur : préfixé d'une apostrophe en CSV, forcé en
texte dans le classeur. Les noms affichés sont saisis par les participants.
"""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any, Iterable

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
# Caractères de contrôle interdits dans le XML d'un classeur.
ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
XLSX_DATE_FORMAT = "dd.mm.yyyy hh:mm"
XLSX_MAX_WIDTH = 60
WIDTH_SAMPLE = 500  # lignes examinées pour régler la largeur des colonnes
README_SHEET = "lisez-moi"


@dataclass
class Table:
    name: str  # nom du fichier CSV / de la feuille (31 caractères au plus)
    description: str
    columns: list[tuple[str, str]]  # (clé, description)
    rows: list[dict[str, Any]] = field(default_factory=list)

    @property
    def keys(self) -> list[str]:
        return [key for key, _ in self.columns]

    def values(self) -> Iterable[list[Any]]:
        for row in self.rows:
            yield [row.get(key) for key in self.keys]


def _dictionary_rows(tables: list[Table]) -> list[list[str]]:
    return [[t.name, key, desc] for t in tables for key, desc in t.columns]


def _readme_lines(tables: list[Table], meta: dict[str, Any]) -> list[str]:
    lines = [f"{key} : {_text(value)}" for key, value in meta.items() if key != "remarques"]
    lines += ["", "Remarques"] + [f"- {note}" for note in meta.get("remarques", [])]
    lines += ["", "Tables"] + [f"- {t.name} ({len(t.rows)} lignes) : {t.description}" for t in tables]
    return lines


def _text(value: Any) -> str:
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    return str(value)


# --- JSON -------------------------------------------------------------------

def _json_default(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    raise TypeError(f"Valeur non exportable : {type(value).__name__}")


def to_json(tables: list[Table], meta: dict[str, Any]) -> bytes:
    payload = {
        "meta": meta,
        "dictionnaire": {
            t.name: {"description": t.description, "colonnes": dict(t.columns)} for t in tables
        },
        "tables": {t.name: [dict(zip(t.keys, values)) for values in t.values()] for t in tables},
    }
    return json.dumps(payload, ensure_ascii=False, indent=1, default=_json_default).encode("utf-8")


# --- CSV --------------------------------------------------------------------

def _csv_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _csv_bytes(header: list[str], rows: Iterable[list[Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows([_csv_cell(v) for v in values] for values in rows)
    return buf.getvalue().encode("utf-8-sig")


def to_csv_zip(tables: list[Table], meta: dict[str, Any]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for t in tables:
            zf.writestr(f"{t.name}.csv", _csv_bytes(t.keys, t.values()))
        zf.writestr(
            "dictionnaire.csv",
            _csv_bytes(["table", "colonne", "description"], _dictionary_rows(tables)),
        )
        zf.writestr("LISEZMOI.txt", "\n".join(_readme_lines(tables, meta)).encode("utf-8-sig"))
    return buf.getvalue()


# --- Excel ------------------------------------------------------------------

def to_xlsx(tables: list[Table], meta: dict[str, Any]) -> bytes:
    # Import local : seul l'export Excel dépend d'openpyxl.
    from openpyxl import Workbook

    wb = Workbook(write_only=True)  # écriture en flux : supporte de gros exports
    for t in tables:
        _write_table_sheet(wb, t)
    _write_readme_sheet(wb, tables, meta)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_table_sheet(wb, table: Table) -> None:
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(table.name)
    rows = list(table.values())
    # En écriture en flux, largeurs, volets et filtre se fixent avant les lignes.
    widths = [len(key) for key in table.keys]
    for values in rows[:WIDTH_SAMPLE]:
        for i, v in enumerate(values):
            if v is not None:
                widths[i] = max(widths[i], min(len(_text(v)), XLSX_MAX_WIDTH))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w + 2
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{len(rows) + 1}"

    ws.append([_bold(ws, key) for key in table.keys])
    for values in rows:
        ws.append([_xlsx_cell(ws, v) for v in values])


def _write_readme_sheet(wb, tables: list[Table], meta: dict[str, Any]) -> None:
    ws = wb.create_sheet(README_SHEET)
    for col, width in (("A", 26), ("B", 30), ("C", 100)):
        ws.column_dimensions[col].width = width
    for key, value in meta.items():
        if key != "remarques":
            ws.append([key, _text(value)])
    ws.append([])
    ws.append([_bold(ws, "Remarques")])
    for note in meta.get("remarques", []):
        ws.append([note])
    ws.append([])
    ws.append([_bold(ws, "table"), _bold(ws, "lignes"), _bold(ws, "description")])
    for t in tables:
        ws.append([t.name, len(t.rows), t.description])
    ws.append([])
    ws.append([_bold(ws, "table"), _bold(ws, "colonne"), _bold(ws, "description")])
    for row in _dictionary_rows(tables):
        ws.append(row)


def _bold(ws, text: str):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    cell = WriteOnlyCell(ws, value=text)
    cell.font = Font(bold=True)
    return cell


def _xlsx_cell(ws, value: Any) -> Any:
    from openpyxl.cell import WriteOnlyCell

    if isinstance(value, datetime):
        cell = WriteOnlyCell(ws, value=value.replace(tzinfo=None))  # Excel ignore les fuseaux
        cell.number_format = XLSX_DATE_FORMAT
        return cell
    if isinstance(value, str):
        value = ILLEGAL_XML_CHARS.sub("", value)
        if value.startswith("="):
            cell = WriteOnlyCell(ws, value=value)
            cell.data_type = "s"  # texte, jamais une formule
            return cell
    return value
