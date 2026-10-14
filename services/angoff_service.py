"""Seuils de niveau par chapitre, fixés par des juges (Angoff modifié).

Proposition v3, § 2.6 (Dwyer et al. 2016) : les seuils provisoires des bandes
(60, 75, 90) sont remplacés, chapitre par chapitre, par le jugement d'enseignant·e·s.

Méthode (variante « réponse la plus probable » de l'Angoff modifié) :
1. pour chaque point de choix, chaque juge indique la réponse que choisirait le
   plus probablement un·e étudiant·e « tout juste compétent·e » de l'année visée ;
2. le seuil d'un·e juge pour un chapitre est la qualité des choix de cet·te
   étudiant·e imaginaire, calculée comme celle d'un joueur :
   Σ (total choisi − pire) / Σ (meilleur − pire) × 100 ;
3. le seuil « Compétent » du chapitre est la médiane des juges, appliquée à partir
   de `MIN_JUDGES` juges ; « À consolider » et « Avancé » restent à `BAND_STEP`
   points en dessous et au-dessus, comme l'écart des seuils provisoires.
"""
from __future__ import annotations

import io
import random
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Mapping

from domain.models import Chapter
from services.scoring_service import BAND_THRESHOLDS, response_scores, total

MIN_JUDGES = 3
BAND_STEP = BAND_THRESHOLDS[1] - BAND_THRESHOLDS[0]  # 15 points entre deux seuils
DEFAULT_SEED = 2026


def point_ref(chapter_id: int, scene_id: int, interaction_id: int) -> str:
    return f"c{chapter_id}-s{scene_id}-i{interaction_id}"


@dataclass
class AngoffPoint:
    ref: str
    chapter_id: int
    chapter: str
    scene: str
    actor: str
    prompt: str
    responses: list[str]
    totals: list[int]


def angoff_points(
    chapters: list[Chapter],
    chapter_ids: Iterable[int] | None = None,
    sample: int | None = None,
    seed: int = DEFAULT_SEED,
) -> list[AngoffPoint]:
    """Points de choix du livre (au moins deux réponses), dans l'ordre du livre."""
    wanted = set(chapter_ids) if chapter_ids is not None else None
    out: list[AngoffPoint] = []
    for ch in chapters:
        if wanted is not None and ch.Id not in wanted:
            continue
        points = [(s, ia) for s in ch.Scenes for ia in s.Interactions if len(ia.Responses) > 1]
        if sample:
            keep = set(random.Random(f"angoff-{seed}-{ch.Id}").sample(range(len(points)), min(sample, len(points))))
            points = [p for i, p in enumerate(points) if i in keep]
        for scene, ia in points:
            out.append(AngoffPoint(
                ref=point_ref(ch.Id, scene.Id, ia.Id),
                chapter_id=ch.Id,
                chapter=ch.Name,
                scene=scene.Title,
                actor=ia.Actor.Name or "",
                prompt=ia.Text or "",
                responses=[r.Text or "" for r in ia.Responses],
                totals=[total(response_scores(r)) for r in ia.Responses],
            ))
    return out


@dataclass
class ParsedAngoff:
    judge: str
    picks: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    unknown: int = 0


def _defined_value(wb, name: str) -> Any:
    if name not in wb.defined_names:
        return None
    for sheet, coord in wb.defined_names[name].destinations:
        if sheet in wb.sheetnames:
            return wb[sheet][coord.replace("$", "")].value
    return None


def parse_angoff_grid(data: bytes, filename: str, points: Mapping[str, AngoffPoint]) -> ParsedAngoff:
    """Lit une grille Angoff remplie : un numéro de réponse par point de choix."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        return ParsedAngoff("", errors=[f"{filename} : ce n'est pas un classeur Excel (.xlsx) lisible."])
    judge = str(_defined_value(wb, "juge") or "").strip() or Path(filename).stem
    parsed = ParsedAngoff(judge[:80])
    if "seuils" not in wb.sheetnames:
        parsed.errors.append(f"{filename} : feuille « seuils » introuvable.")
        return parsed
    rows = wb["seuils"].iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows, [])]
    col = {h: i for i, h in enumerate(header)}
    if "point" not in col or "réponse la plus probable (n°)" not in col:
        parsed.errors.append(f"{filename} : colonnes « point » ou « réponse la plus probable (n°) » manquantes.")
        return parsed
    for line, row in enumerate(rows, start=2):
        ref = str(row[col["point"]] or "").strip()
        raw = row[col["réponse la plus probable (n°)"]]
        if not ref or raw is None or (isinstance(raw, str) and not raw.strip()):
            continue
        point = points.get(ref)
        if point is None:
            parsed.unknown += 1
            continue
        try:
            no = int(float(str(raw).strip()))
        except ValueError:
            no = 0
        if not 1 <= no <= len(point.responses):
            parsed.errors.append(f"{filename}, ligne {line} : « {raw} » n'est pas un numéro de réponse de 1 à {len(point.responses)}.")
            continue
        comment = str(row[col["commentaire"]] or "").strip() if "commentaire" in col else ""
        parsed.picks.append({"point": ref, "chapter_id": point.chapter_id, "response_no": no, "comment": comment})
    return parsed


def judge_cut(points: list[AngoffPoint], picks: Mapping[str, int]) -> float | None:
    """Qualité des choix de l'étudiant·e « tout juste compétent·e » d'un·e juge."""
    gain = span = 0.0
    for p in points:
        no = picks.get(p.ref)
        lo, hi = min(p.totals), max(p.totals)
        if no is None or hi <= lo:
            continue
        gain += p.totals[no - 1] - lo
        span += hi - lo
    return round(100 * gain / span, 1) if span else None


def cut_scores(points: list[AngoffPoint], ratings: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    """Seuil de chaque chapitre : un par juge, médiane, et seuils de bande appliqués."""
    by_judge: dict[str, dict[str, int]] = {}
    for r in ratings:
        by_judge.setdefault(r["judge"], {})[r["point"]] = r["response_no"]
    chapters: dict[int, list[AngoffPoint]] = {}
    names: dict[int, str] = {}
    for p in points:
        chapters.setdefault(p.chapter_id, []).append(p)
        names[p.chapter_id] = p.chapter
    out = []
    for cid, pts in chapters.items():
        refs = {p.ref for p in pts}
        cuts = {
            judge: cut
            for judge, picks in by_judge.items()
            if (cut := judge_cut(pts, picks)) is not None
        }
        judged = [sum(1 for ref in picks if ref in refs) for picks in by_judge.values()]
        median = round(statistics.median(cuts.values()), 1) if cuts else None
        applied = median is not None and len(cuts) >= MIN_JUDGES
        out.append({
            "chapter_id": cid,
            "chapter": names[cid],
            "n_points": len(pts),
            "n_judged": max(judged, default=0),
            "cuts": cuts,
            "n_judges": len(cuts),
            "median": median,
            "min": min(cuts.values()) if cuts else None,
            "max": max(cuts.values()) if cuts else None,
            "applied": applied,
            "thresholds": thresholds_from(median) if applied else None,
        })
    return out


def thresholds_from(cut: float) -> tuple[float, float, float]:
    """Seuils À consolider, Compétent, Avancé autour du seuil « Compétent »."""
    return (max(0.0, round(cut - BAND_STEP, 1)), cut, min(100.0, round(cut + BAND_STEP, 1)))


def chapter_thresholds(cuts: list[dict[str, Any]]) -> dict[int, tuple[float, float, float]]:
    """Seuils appliqués (au moins `MIN_JUDGES` juges), par chapitre."""
    return {c["chapter_id"]: c["thresholds"] for c in cuts if c["applied"]}
