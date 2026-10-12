"""Validation de la clé de cotation par un panel d'experts (scoring v2, étape 2).

Démarche (docs/Scoring_v2_recherche_et_plan.md, § 3.1) :
1. chaque expert reçoit une grille Excel des réponses aux points de choix,
   **sans les points de l'auteur** (cotation à l'aveugle) ;
2. il cote chaque réponse de −3 à +3 sur les 7 compétences, la marque
   « critique » au besoin, et la justifie en une phrase ;
3. les grilles remplies sont importées ; on calcule l'accord entre experts
   (ICC à deux facteurs, effets aléatoires, accord absolu : ICC(2,1) pour un
   expert, ICC(2,k) pour la moyenne des k experts ; Koo & Li 2016) et on
   repère les réponses à revoir au tour suivant (méthode Delphi).

La clé consensuelle (médiane des experts) n'est pas encore appliquée au livre.
"""
from __future__ import annotations

import io
import random
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Mapping

from domain.models import Chapter
from domain.soft_skills import SOFT_SKILL_LABELS, SOFT_SKILLS

SCORE_MIN, SCORE_MAX = -3, 3
# Réponse marquée critique à titre provisoire, en attendant les experts.
CRITICAL_PROVISIONAL_TOTAL = -7
# Réponse à revoir : experts trop éloignés entre eux, ou de la clé de l'auteur.
DISAGREEMENT_RANGE = 3
AUTHOR_GAP = 2
DEFAULT_SEED = 2026
TOTAL = "Total"


def response_ref(chapter_id: int, scene_id: int, interaction_id: int, response_no: int) -> str:
    """Identifiant stable d'une réponse : rang commençant à 1 dans l'ordre du livre."""
    return f"c{chapter_id}-s{scene_id}-i{interaction_id}-r{response_no}"


@dataclass
class KeyItem:
    """Une réponse proposée à un point de choix (au moins deux réponses)."""

    ref: str
    chapter_id: int
    chapter: str
    scene_id: int
    scene: str
    interaction_id: int
    actor: str
    prompt: str
    response_no: int
    text: str
    author: dict[str, int]
    critical: bool

    @property
    def author_total(self) -> int:
        return sum(self.author.values())

    @property
    def provisional_critical(self) -> bool:
        return self.critical or self.author_total <= CRITICAL_PROVISIONAL_TOTAL


def key_items(
    chapters: list[Chapter],
    chapter_ids: Iterable[int] | None = None,
    sample: int | None = None,
    seed: int = DEFAULT_SEED,
) -> list[KeyItem]:
    """Réponses des points de choix, dans l'ordre du livre.

    `sample` : nombre de points de choix tirés au hasard par chapitre (graine
    fixe : le même tirage d'un export à l'autre, pour tous les experts).
    """
    wanted = set(chapter_ids) if chapter_ids is not None else None
    items: list[KeyItem] = []
    for ch in chapters:
        if wanted is not None and ch.Id not in wanted:
            continue
        points = [(s, ia) for s in ch.Scenes for ia in s.Interactions if len(ia.Responses) > 1]
        if sample:
            keep = set(random.Random(f"{seed}-{ch.Id}").sample(range(len(points)), min(sample, len(points))))
            points = [p for i, p in enumerate(points) if i in keep]
        for scene, ia in points:
            for no, r in enumerate(ia.Responses, start=1):
                items.append(KeyItem(
                    ref=response_ref(ch.Id, scene.Id, ia.Id, no),
                    chapter_id=ch.Id,
                    chapter=ch.Name,
                    scene_id=scene.Id,
                    scene=scene.Title,
                    interaction_id=ia.Id,
                    actor=ia.Actor.Name or "",
                    prompt=ia.Text or "",
                    response_no=no,
                    text=r.Text or "",
                    author={k: int(r.score(k) or 0) for k in SOFT_SKILLS},
                    critical=bool(r._extra.get("Critical")),
                ))
    return items


# --- Lecture d'une grille remplie -------------------------------------------

@dataclass
class ParsedGrid:
    expert: str
    round: int
    ratings: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    unknown_refs: int = 0  # réponses cotées qui n'existent plus dans le livre


_YES = {"oui", "o", "x", "1", "true", "vrai", "yes", "critique"}


def _score(value: Any) -> int | None:
    """None si vide ; ValueError si ce n'est pas un entier de −3 à +3."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError(value)
    if isinstance(value, str):
        value = value.strip().replace("−", "-").replace("+", "")
    number = float(value)
    if number != int(number) or not SCORE_MIN <= number <= SCORE_MAX:
        raise ValueError(value)
    return int(number)


def _defined_value(wb, name: str) -> Any:
    if name not in wb.defined_names:
        return None
    for sheet, coord in wb.defined_names[name].destinations:
        if sheet in wb.sheetnames:
            return wb[sheet][coord.replace("$", "")].value
    return None


def parse_grid(data: bytes, filename: str, known_refs: set[str]) -> ParsedGrid:
    """Lit une grille d'expert remplie. Les cases vides d'une réponse cotée valent 0."""
    from openpyxl import load_workbook  # import local : seule l'importation en dépend

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:  # fichier illisible, pas un .xlsx
        return ParsedGrid("", 1, errors=[f"{filename} : ce n'est pas un classeur Excel (.xlsx) lisible."])
    expert = str(_defined_value(wb, "expert") or "").strip() or Path(filename).stem
    try:
        round_ = int(_defined_value(wb, "tour") or 1)
    except (TypeError, ValueError):
        round_ = 1
    grid = ParsedGrid(expert[:80], round_)
    if "cotation" not in wb.sheetnames:
        grid.errors.append(f"{filename} : feuille « cotation » introuvable.")
        return grid

    rows = wb["cotation"].iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows, [])]
    col = {h: i for i, h in enumerate(header)}
    needed = ["ref", *(SOFT_SKILL_LABELS[k] for k in SOFT_SKILLS), "critique", "justification"]
    missing = [h for h in needed if h not in col]
    if missing:
        grid.errors.append(f"{filename} : colonnes manquantes : {', '.join(missing)}.")
        return grid

    for line, row in enumerate(rows, start=2):
        ref = str(row[col["ref"]] or "").strip()
        if not ref:
            continue
        scores: dict[str, int | None] = {}
        for k in SOFT_SKILLS:
            raw = row[col[SOFT_SKILL_LABELS[k]]]
            try:
                scores[k] = _score(raw)
            except (TypeError, ValueError):
                grid.errors.append(
                    f"{filename}, ligne {line} ({SOFT_SKILL_LABELS[k]}) : « {raw} » n'est pas un "
                    f"entier de {SCORE_MIN} à +{SCORE_MAX}."
                )
                scores[k] = None
        critical = str(row[col["critique"]] or "").strip().lower() in _YES
        rationale = str(row[col["justification"]] or "").strip()
        comment = str(row[col["commentaire"]] or "").strip() if "commentaire" in col else ""
        if all(v is None for v in scores.values()) and not critical:
            continue  # réponse pas encore cotée
        if ref not in known_refs:  # réponse d'une autre version du livre
            grid.unknown_refs += 1
            continue
        grid.ratings.append({
            "ref": ref,
            "scores": {k: v or 0 for k, v in scores.items()},
            "critical": critical,
            "rationale": rationale,
            "comment": comment,
        })
    return grid


# --- Accord entre experts ---------------------------------------------------

def icc(matrix: list[list[float]]) -> tuple[float | None, float | None]:
    """ICC(2,1) et ICC(2,k) : deux facteurs, effets aléatoires, accord absolu.

    `matrix` : une ligne par réponse, une colonne par expert, sans case vide.
    (Shrout & Fleiss 1979 ; McGraw & Wong 1996 ; Koo & Li 2016.)
    """
    n = len(matrix)
    k = len(matrix[0]) if matrix else 0
    if n < 2 or k < 2:
        return None, None
    grand = sum(map(sum, matrix)) / (n * k)
    rows = [sum(r) / k for r in matrix]
    cols = [sum(r[j] for r in matrix) / n for j in range(k)]
    ss_rows = k * sum((m - grand) ** 2 for m in rows)
    ss_cols = n * sum((m - grand) ** 2 for m in cols)
    ss_total = sum((x - grand) ** 2 for r in matrix for x in r)
    ss_error = ss_total - ss_rows - ss_cols
    msr = ss_rows / (n - 1)
    msc = ss_cols / (k - 1)
    mse = ss_error / ((n - 1) * (k - 1))
    single = msr + (k - 1) * mse + k * (msc - mse) / n
    average = msr + (msc - mse) / n
    if single <= 0 or average <= 0:
        return None, None
    return round((msr - mse) / single, 3), round((msr - mse) / average, 3)


def icc_label(value: float | None) -> str:
    """Interprétation de Koo & Li (2016)."""
    if value is None:
        return "non calculable"
    if value < 0.5:
        return "faible"
    if value < 0.75:
        return "modéré"
    if value < 0.9:
        return "bon"
    return "excellent"


def _total(scores: Mapping[str, int]) -> int:
    return sum(int(scores.get(k, 0) or 0) for k in SOFT_SKILLS)


def analyse(items: list[KeyItem], ratings: Iterable[Mapping[str, Any]]) -> dict[str, Any]:
    """Accord entre experts et réponses à revoir, pour un tour.

    `ratings` : {ref, expert, scores, critical, rationale, comment}.
    L'ICC se calcule sur les réponses cotées par tous les experts du tour.
    """
    by_ref: dict[str, dict[str, Mapping[str, Any]]] = {}
    for r in ratings:
        by_ref.setdefault(r["ref"], {})[r["expert"]] = r
    experts = sorted({e for rated in by_ref.values() for e in rated})
    complete = [it for it in items if len(by_ref.get(it.ref, {})) == len(experts)]

    agreement = []
    for k in [*SOFT_SKILLS, TOTAL]:
        def value(r: Mapping[str, Any]) -> float:
            return _total(r["scores"]) if k == TOTAL else r["scores"].get(k, 0)

        matrix = [[value(by_ref[it.ref][e]) for e in experts] for it in complete]
        single, average = icc(matrix)
        agreement.append({
            "skill": k,
            "label": SOFT_SKILL_LABELS.get(k, "Total des 7 compétences"),
            "n_items": len(matrix),
            "icc_single": single,
            "icc_average": average,
            "level_single": icc_label(single),
            "level_average": icc_label(average),
        })

    rows = []
    for it in items:
        rated = by_ref.get(it.ref)
        if not rated:
            continue
        stats: dict[str, dict[str, float]] = {}
        for k in SOFT_SKILLS:
            values = [r["scores"].get(k, 0) for r in rated.values()]
            stats[k] = {"median": statistics.median(values), "min": min(values), "max": max(values)}
        totals = [_total(r["scores"]) for r in rated.values()]
        disagree = [k for k in SOFT_SKILLS if stats[k]["max"] - stats[k]["min"] >= DISAGREEMENT_RANGE]
        gap = [k for k in SOFT_SKILLS if abs(stats[k]["median"] - it.author[k]) >= AUTHOR_GAP]
        votes = sum(1 for r in rated.values() if r.get("critical"))
        rows.append({
            "item": it,
            "n_experts": len(rated),
            "stats": stats,
            "median_total": statistics.median(totals),
            "disagreement": disagree,
            "author_gap": gap,
            "critical_votes": votes,
            "critical_majority": votes * 2 > len(rated),
            "rationales": [f"{e} : {r['rationale']}" for e, r in sorted(rated.items()) if r.get("rationale")],
            "comments": [f"{e} : {r['comment']}" for e, r in sorted(rated.items()) if r.get("comment")],
            "to_review": bool(disagree or gap),
        })
    return {
        "experts": experts,
        "n_rated": len(rows),
        "n_complete": len(complete),
        "agreement": agreement,
        "rows": rows,
        "n_to_review": sum(1 for r in rows if r["to_review"]),
        "n_critical": sum(1 for r in rows if r["critical_majority"]),
    }
