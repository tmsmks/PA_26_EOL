"""Validation de la clé de cotation : grille d'experts, import, accord (ICC), synthèse."""
import io
import json
import re
import zipfile

import pytest
from openpyxl import load_workbook

from app import create_app
from domain.models import Chapter
from domain.soft_skills import SOFT_SKILL_LABELS, SOFT_SKILLS
from exporters.expert_key import grid_workbook
from repositories.orm import ROLE_TRAINER
from repositories.users import UserRepository
from services.auth_service import hash_password
from services.expert_key_service import analyse, icc, icc_label, key_items, parse_grid


def _scene(critical=False):
    """1 → (EM+2 → 2 | EM-1 → fin) ; 2 → (RD+1 → fin | RD-3 PB-3 CC-2 → fin, critique ?)."""
    def resp(next_id, extra=None, **scores):
        return {"Id": 1, "Text": f"vers {next_id}", "NextInteractionID": next_id,
                "SoftSkillDimensions": scores, **(extra or {})}

    return {
        "Id": 1,
        "Title": "Scène",
        "Interactions": [
            {"Id": 1, "Text": "Bonjour", "Actor": {"Name": "Mme A"}, "Responses": [
                resp(2, Empathy=2), resp(-1, Empathy=-1)]},
            {"Id": 2, "Text": "Et alors ?", "Actor": {"Name": "Mme A"}, "Responses": [
                resp(-1, RespectAndDignity=1),
                resp(-1, {"Critical": True} if critical else None,
                     RespectAndDignity=-3, ProfessionalBoundaries=-3, CommunicationClarity=-2)]},
            {"Id": 3, "Text": "Fin.", "Responses": [resp(-1)]},
        ],
    }


BOOK = {"Chapters": [
    {"Id": 1, "Name": "Un", "Intro": "", "Scenes": [_scene(critical=True)]},
    {"Id": 2, "Name": "Deux", "Intro": "", "Scenes": [_scene()]},
]}


def _chapters():
    return [Chapter.from_dict(ch) for ch in BOOK["Chapters"]]


def test_icc_matches_shrout_and_fleiss():
    """Exemple publié de Shrout & Fleiss (1979) : 6 cibles, 4 juges."""
    data = [[9, 2, 5, 8], [6, 1, 3, 2], [8, 4, 6, 8], [7, 1, 2, 6], [10, 5, 6, 9], [6, 2, 4, 7]]
    single, average = icc(data)
    assert round(single, 2) == 0.29 and round(average, 2) == 0.62
    assert icc([[1, 1]]) == (None, None)
    assert [icc_label(v) for v in (0.3, 0.6, 0.8, 0.95, None)] == [
        "faible", "modéré", "bon", "excellent", "non calculable"]


def test_key_items_list_choice_points_only():
    items = key_items(_chapters())
    assert [it.ref for it in items[:4]] == ["c1-s1-i1-r1", "c1-s1-i1-r2", "c1-s1-i2-r1", "c1-s1-i2-r2"]
    assert len(items) == 8  # l'interaction 3 (« Continuer ») n'est pas un point de choix
    critical = items[3]
    assert critical.author["ProfessionalBoundaries"] == -3 and critical.author_total == -8
    assert critical.critical and critical.provisional_critical
    assert not items[7].critical and items[7].provisional_critical  # total ≤ −7
    assert {it.chapter_id for it in key_items(_chapters(), [2])} == {2}
    sample = key_items(_chapters(), sample=1)
    assert len({(it.chapter_id, it.interaction_id) for it in sample}) == 2  # 1 point par chapitre
    assert [it.ref for it in sample] == [it.ref for it in key_items(_chapters(), sample=1)]


def _fill(data: bytes, expert: str, values: dict, critical: set = frozenset(), bad: str | None = None) -> bytes:
    """Remplit une grille comme le ferait un expert. values : ref → {compétence: cote}."""
    wb = load_workbook(io.BytesIO(data))
    wb["consignes"]["B4"] = expert
    ws = wb["cotation"]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        ref = row[0].value
        for k, v in values.get(ref, {}).items():
            row[header.index(SOFT_SKILL_LABELS[k])].value = v
        if ref in critical:
            row[header.index("critique")].value = "oui"
            row[header.index("justification")].value = "Viole le secret."
        if ref == bad:
            row[header.index(SOFT_SKILL_LABELS["Empathy"])].value = 7
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_grid_is_blind_and_round_trips():
    items = key_items(_chapters())
    data = grid_workbook("Livre test", items, 2, "Tous les chapitres")
    wb = load_workbook(io.BytesIO(data))
    ws = wb["cotation"]
    header = [c.value for c in ws[1]]
    assert header[:7] == ["ref", "chapitre", "scène", "personnage", "réplique", "réponse n°", "réponse"]
    assert ws.max_row == len(items) + 1
    # À l'aveugle : aucune cote de l'auteur dans la grille.
    scores = [c.value for row in ws.iter_rows(min_row=2, min_col=8, max_col=14) for c in row]
    assert set(scores) == {None}
    assert wb["consignes"]["B5"].value == 2  # tour

    known = {it.ref for it in items}
    filled = _fill(data, "Dre X", {"c1-s1-i1-r1": {"Empathy": 3}, "c1-s1-i2-r2": {"ProfessionalBoundaries": -3}},
                   critical={"c1-s1-i2-r2"})
    grid = parse_grid(filled, "grille.xlsx", known)
    assert (grid.expert, grid.round, grid.errors, grid.unknown_refs) == ("Dre X", 2, [], 0)
    assert [r["ref"] for r in grid.ratings] == ["c1-s1-i1-r1", "c1-s1-i2-r2"]  # lignes vides ignorées
    first = grid.ratings[0]
    assert first["scores"]["Empathy"] == 3 and first["scores"]["Compassion"] == 0  # vide = 0
    assert grid.ratings[1]["critical"] and grid.ratings[1]["rationale"] == "Viole le secret."

    bad = parse_grid(_fill(data, "Dre X", {}, bad="c1-s1-i1-r2"), "g.xlsx", known)
    assert len(bad.errors) == 1 and "ligne 3" in bad.errors[0]
    assert parse_grid(filled, "g.xlsx", {"c1-s1-i1-r1"}).unknown_refs == 1
    assert parse_grid(b"pas un classeur", "g.xlsx", known).errors


def test_analyse_flags_disagreements_author_gaps_and_critical_majority():
    items = key_items(_chapters(), [1])
    ref = "c1-s1-i2-r2"  # auteur : RD −3, PB −3, CC −2
    zero = {k: 0 for k in SOFT_SKILLS}
    ratings = [
        {"ref": ref, "expert": "A", "scores": {**zero, "ProfessionalBoundaries": -3}, "critical": True, "rationale": "secret"},
        {"ref": ref, "expert": "B", "scores": {**zero, "ProfessionalBoundaries": 0}, "critical": True, "rationale": ""},
        {"ref": ref, "expert": "C", "scores": {**zero, "ProfessionalBoundaries": -2}, "critical": False, "rationale": ""},
    ]
    result = analyse(items, ratings)
    row = result["rows"][0]
    assert result["experts"] == ["A", "B", "C"] and result["n_rated"] == 1
    assert row["disagreement"] == ["ProfessionalBoundaries"]  # −3 à 0 : écart de 3
    assert set(row["author_gap"]) == {"RespectAndDignity", "CommunicationClarity"}  # médiane 0
    assert row["critical_votes"] == 2 and row["critical_majority"]
    assert row["rationales"] == ["A : secret"]
    assert result["agreement"][-1]["skill"] == "Total"
    assert result["agreement"][-1]["icc_single"] is None  # une seule réponse : pas d'ICC


@pytest.fixture()
def app(tmp_path):
    (tmp_path / "chapter.json").write_text(json.dumps(BOOK), encoding="utf-8")
    app = create_app({
        "TESTING": True,
        "SECRET_KEY": "test",
        "DATABASE_URL": "sqlite:///" + str(tmp_path / "test.db"),
        "CHAPTERS_BASE_DIR": str(tmp_path),
        "CSRF_ENABLED": False,
    })
    with app.app_context():
        from repositories.database import db_session
        UserRepository(db_session()).create("trainer", "Formatrice", hash_password("motdepasse1"), ROLE_TRAINER)
    return app


def _trainer(app):
    t = app.test_client()
    t.post("/login", data={"username": "trainer", "password": "motdepasse1"})
    return t


def test_trainer_page_downloads_grids_imports_and_summarizes(app):
    alice = app.test_client()
    alice.post("/register", data={"username": "alice", "display_name": "Alice",
                                  "password": "motdepasse1", "password_confirm": "motdepasse1"})
    assert alice.get("/trainer/cle/").status_code == 403

    t = _trainer(app)
    page = t.get("/trainer/cle/?book=end_of_life").get_data(as_text=True)
    assert "Clé de cotation" in page and "8 réponses aux points de choix" in page
    r = t.get("/trainer/cle/end_of_life/grille?chapter=1&round=1")
    assert r.status_code == 200 and "grille_experts_end_of_life_ch1_tour1.xlsx" in r.headers["Content-Disposition"]
    assert t.get("/trainer/cle/end_of_life/grille?sample=7").status_code == 400
    assert t.get("/trainer/cle/end_of_life/grille?chapter=9").status_code == 404

    grid = t.get("/trainer/cle/end_of_life/grille").data
    refs = [it.ref for it in key_items(_chapters())]
    files = []
    for name, shift in (("A", 0), ("B", 1), ("C", -1)):
        values = {ref: {"Empathy": max(-3, min(3, i % 4 - 1 + shift))} for i, ref in enumerate(refs)}
        files.append((io.BytesIO(_fill(grid, name, values)), f"{name}.xlsx"))
    r = t.post("/trainer/cle/end_of_life/import", data={"grids": files},
               content_type="multipart/form-data", follow_redirects=True)
    page = r.get_data(as_text=True)
    assert page.count("8 réponse(s) cotée(s) importée(s)") == 3
    assert "Accord entre experts" in page and "ICC(2,1)" in page

    synth = t.get("/trainer/cle/end_of_life/synthese?round=1")
    wb = load_workbook(io.BytesIO(synth.data))
    assert wb.sheetnames == ["accord", "reponses", "lisez-moi"]
    accord = {row[0]: row for row in wb["accord"].iter_rows(min_row=2, values_only=True)}
    assert accord["Empathie"][1] == 8 and accord["Empathie"][3] is not None

    t.post("/trainer/cle/end_of_life/supprimer", data={"expert": "C", "round": 1})
    page = t.get("/trainer/cle/?book=end_of_life").get_data(as_text=True)
    assert re.search(r"<td>C</td>", page) is None and "<td>A</td>" in page


def test_critical_answers_reach_the_export(app):
    alice = app.test_client()
    alice.post("/register", data={"username": "alice", "display_name": "Alice",
                                  "password": "motdepasse1", "password_confirm": "motdepasse1"})
    sess = alice.post("/api/play/end_of_life/session").get_json()["session"]
    alice.post(sess["urls"]["attempts"], json={"chapter_id": 1, "scene_id": 1, "entry_id": 1, "choices": [0, 1]})
    data = json.loads(_trainer(app).get("/trainer/export?format=json").data)
    choix = data["tables"]["choix"]
    assert [(c["critique"], c["critique_provisoire"]) for c in choix] == [(False, False), (True, True)]
    zf = zipfile.ZipFile(io.BytesIO(_trainer(app).get("/trainer/export?format=csv").data))
    assert "critique" in zf.read("choix.csv").decode("utf-8-sig").splitlines()[0]



# --- Relecture ciblée et seuils de niveau (Angoff) ---------------------------

def test_targeted_review_keeps_disputed_best_answers_and_critical_ones(tmp_path):
    from services.expert_key_service import KEY_FILES, targeted_review

    key = {"reponses": {
        # Point 1 : la meilleure réponse a changé (r1 → r2).
        "c1-s1-i1-r1": {"cotes": "rd2", "auteur": "rd3", "critique": False},
        "c1-s1-i1-r2": {"cotes": "rd3", "auteur": "rd1", "critique": False},
        "c1-s1-i1-r3": {"cotes": "rd-2", "auteur": "rd-2", "critique": True},
        # Point 2 : même meilleure réponse.
        "c1-s1-i2-r1": {"cotes": "em2", "auteur": "em3", "critique": False},
        "c1-s1-i2-r2": {"cotes": "em-1", "auteur": "em-1", "critique": False},
    }}
    (tmp_path / KEY_FILES["medstudents_y2"]).write_text(json.dumps(key), encoding="utf-8")
    t = targeted_review("medstudents_y2", tmp_path)
    assert t["refs"] == {"c1-s1-i1-r1", "c1-s1-i1-r2", "c1-s1-i1-r3"}
    assert t["n_points_changed"] == 1 and t["n_critical"] == 1 and t["n_items"] == 3
    assert targeted_review("end_of_life", tmp_path) is None
    assert key_items(_chapters(), refs={"c1-s1-i1-r2"})[0].ref == "c1-s1-i1-r2"


def test_angoff_cut_scores_apply_from_three_judges():
    from services.angoff_service import angoff_points, cut_scores, thresholds_from

    points = angoff_points(_chapters())
    assert [p.ref for p in points] == ["c1-s1-i1", "c1-s1-i2", "c2-s1-i1", "c2-s1-i2"]

    def picks(judge, first, second):
        return [{"judge": judge, "point": "c1-s1-i1", "response_no": first},
                {"judge": judge, "point": "c1-s1-i2", "response_no": second}]

    # Meilleure réponse puis pire : (3 + 0) / (3 + 9) = 25 ; deux meilleures : 100.
    two = cut_scores(points, picks("A", 1, 2) + picks("B", 1, 1))
    ch1 = next(c for c in two if c["chapter_id"] == 1)
    assert ch1["cuts"] == {"A": 25.0, "B": 100.0} and not ch1["applied"] and ch1["thresholds"] is None
    three = cut_scores(points, picks("A", 1, 2) + picks("B", 1, 1) + picks("C", 1, 1))
    ch1 = next(c for c in three if c["chapter_id"] == 1)
    assert ch1["median"] == 100.0 and ch1["applied"]
    assert ch1["thresholds"] == thresholds_from(100.0) == (85.0, 100.0, 100.0)


def test_angoff_grid_round_trips_and_sets_the_chapter_bands(app):
    t = _trainer(app)
    page = t.get("/trainer/cle/?book=end_of_life").get_data(as_text=True)
    assert "Seuils de niveau" in page and "les seuils provisoires s'appliquent" in page
    r = t.get("/trainer/cle/end_of_life/angoff/grille")
    assert r.status_code == 200 and "grille_seuils_end_of_life_tout.xlsx" in r.headers["Content-Disposition"]

    files = []
    for judge, second in (("A", 2), ("B", 1), ("C", 1)):
        wb = load_workbook(io.BytesIO(r.data))
        wb["consignes"]["B4"] = judge
        ws = wb["seuils"]
        header = [c.value for c in ws[1]]
        col = header.index("réponse la plus probable (n°)") + 1
        for row in range(2, ws.max_row + 1):
            point = ws.cell(row=row, column=1).value
            ws.cell(row=row, column=col, value=second if point.endswith("i2") else 1)
        buf = io.BytesIO()
        wb.save(buf)
        files.append((io.BytesIO(buf.getvalue()), f"{judge}.xlsx"))
    r = t.post("/trainer/cle/end_of_life/angoff/import", data={"grids": files},
               content_type="multipart/form-data", follow_redirects=True)
    page = r.get_data(as_text=True)
    assert "A : 4 point(s) de choix jugé(s) importé(s)" in page
    assert "85,0 · 100,0 · 100,0" in page  # médiane des 3 juges = 100

    with app.test_request_context():
        from views.scoring_context import book_thresholds
        assert book_thresholds("end_of_life", _chapters()) == {1: (85.0, 100.0, 100.0), 2: (85.0, 100.0, 100.0)}

    bad = load_workbook(io.BytesIO(t.get("/trainer/cle/end_of_life/angoff/grille").data))
    bad["seuils"].cell(row=2, column=7, value=9)
    buf = io.BytesIO()
    bad.save(buf)
    page = t.post("/trainer/cle/end_of_life/angoff/import", data={"grids": [(io.BytesIO(buf.getvalue()), "Z.xlsx")]},
                  content_type="multipart/form-data", follow_redirects=True).get_data(as_text=True)
    assert "n'est pas un numéro de réponse" in page and "grille non importée" in page
